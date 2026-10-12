# -*- coding: utf-8 -*-
"""
核对报告输出（§8.2）：核对报告.xlsx + 核对报告.md。
开头打印列定位结果；逐样本一行；「金额命中·摘要不达标」专项标注。
"""
from __future__ import annotations
from typing import List, Dict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

COLUMNS = ["样本序号", "明细账金额", "样本摘要", "识别金额(命中口径)", "金额差额",
           "发票金额", "付款单金额", "金额一致性", "摘要相似度", "合同号", "发票号", "付款单号",
           "是否跨期", "证据源追溯", "OCR置信度", "来源引擎", "判定", "原因/备注"]

ROW_KEYS = ["sid", "amount", "summary", "hit_amount", "delta",
            "inv_amt", "pay_amt", "amt_consistency", "sim", "contract", "invoice", "payment",
            "cross_period", "evidence_trace", "ocr_conf", "source_engine", "state", "reason"]

FILL_OK = PatternFill("solid", fgColor="D7F0D5")
FILL_WARN = PatternFill("solid", fgColor="FBF2CC")
FILL_BAD = PatternFill("solid", fgColor="F6D6D2")
FILL_HEAD = PatternFill("solid", fgColor="2C5AA0")


def _fill_for(state: str):
    if state.startswith("✅"):
        return FILL_OK
    if state.startswith("⚠"):
        return FILL_WARN
    return FILL_BAD


def write_xlsx(path, layout, rows: List[Dict], engine_status: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "核对报告"

    # 顶部：列定位结果
    ws.append(["列定位结果（按表头文字动态命中）"])
    ws.append([f"工作表：{layout.sheet}　表头行：{layout.header_row}"])
    for f, ch in layout.columns.items():
        tag = "（模糊命中，请核对）" if ch.fuzzy else ""
        ws.append([f"  {f}", f"{ch.col_letter} 列", ch.actual_header, f"相似度{ch.score:.2f}{tag}"])
    ws.append([f"引擎状态：{engine_status}"])
    for w in layout.warnings:
        ws.append([f"⚠ {w}"])
    ws.append([])

    head_idx = ws.max_row + 1
    ws.append(COLUMNS)
    for c in ws[head_idx]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")

    state_i = COLUMNS.index("判定") + 1
    reason_i = COLUMNS.index("原因/备注") + 1
    for r in rows:
        ws.append([r.get(k, "") for k in ROW_KEYS])
        last = ws.max_row
        ws.cell(last, state_i).fill = _fill_for(r["state"])
        if r.get("amount_flag"):
            ws.cell(last, reason_i).font = Font(color="C0392B", bold=True)
        for c in ws[last]:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    widths = [9, 13, 24, 17, 9, 16, 16, 20, 9, 18, 20, 24, 30, 36, 9, 14, 9, 36]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)


def write_md(path, layout, rows: List[Dict], engine_status: str, safety: str):
    lines = []
    lines.append("# TOD 核对报告\n")
    lines.append("> " + safety.replace("\n", "\n> ") + "\n")
    lines.append("## 列定位结果（动态表头）\n")
    lines.append(f"- 工作表：`{layout.sheet}`　表头行：**{layout.header_row}**")
    for f, ch in layout.columns.items():
        tag = " ⚠模糊命中" if ch.fuzzy else ""
        lines.append(f"- {f}：**{ch.col_letter}** 列　『{ch.actual_header}』　相似度{ch.score:.2f}{tag}")
    lines.append(f"- 引擎状态：{engine_status}")
    for w in layout.warnings:
        lines.append(f"- ⚠ {w}")
    lines.append("")

    lines.append("## 逐样本核对结果\n")
    lines.append("| " + " | ".join(COLUMNS) + " |")
    lines.append("|" + "|".join(["---"] * len(COLUMNS)) + "|")
    for r in rows:
        cells = [str(r.get(k, "")) for k in ROW_KEYS]
        cells = [c.replace("\n", "<br>").replace("|", "/") for c in cells]
        lines.append("| " + " | ".join(cells) + " |")
    lines.append("")

    incons = [r for r in rows if "⚠" in str(r.get("amt_consistency", ""))]
    if incons:
        lines.append("## ⚠ 发票/付款单金额不一致（需人工核对凭证）\n")
        for r in incons:
            lines.append(f"- 样本 **{r['sid']}**：发票金额[{r['inv_amt']}] 付款单金额[{r['pay_amt']}] —— {r['amt_consistency']}")
        lines.append("")

    flagged = [r for r in rows if r.get("amount_flag")]
    if flagged:
        lines.append("## ⚠ 金额命中·摘要不达标（优先人工确认摘要）\n")
        for r in flagged:
            lines.append(f"- 样本 **{r['sid']}**：{r['reason']}")
        lines.append("")

    n_ok = sum(1 for r in rows if r["state"].startswith("✅"))
    n_warn = sum(1 for r in rows if r["state"].startswith("⚠"))
    n_bad = sum(1 for r in rows if r["state"].startswith("❌"))
    lines.append(f"## 汇总\n\n- ✅ 一致 {n_ok}　⚠️ 存疑 {n_warn}　❌ 不一致 {n_bad}　合计 {len(rows)}")
    lines.append("- 仅 ✅ 一致样本会自动填入底稿目标列；⚠️/❌ 留空待人工。请先看本报告再使用已填底稿。")

    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
